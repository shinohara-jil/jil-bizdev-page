# コーポレートサイト企画書（261029_コーポレートサイト企画書_v0.3.md）を PowerPoint にする
# JILデザインの共通部品（design-jil）を使い、色は新しいコーポレートサイトに合わせて黄色→青にしている
import sys
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
sys.path.insert(0, str(Path.home() / '.claude' / 'skills' / 'design-jil'))
from jil_deck import *
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_CONNECTOR

HERE = Path(__file__).parent
IMG = HERE / '_img'
OUT = HERE / '261029_コーポレートサイト.pptx'

# サイトの色（docs/index.html の --ink / --accent / --mute / --paper）
INK2 = RGBColor(0x0E, 0x10, 0x13)
ACC = RGBColor(0x2D, 0x44, 0xE0)
MUTE = RGBColor(0x5F, 0x66, 0x70)
PAPER = RGBColor(0xF1, 0xF2, 0xF4)
Y0 = 2.35  # 本文の開始位置


def slide(prs, no, title, lead):
    """左上「● 01 ── 欄の名前」（サイトの欄見出しと同じ形）、右上に JIL. のロゴ、大きな見出し"""
    s = prs.slides.add_slide(prs._jil_layout)
    rect(s, L, 0.61, 0.08, 0.08, fill=INK2, shape=MSO_SHAPE.OVAL)
    tb(s, L + 0.2, 0.5, 0.4, 0.3, f'{no:02d}', size=13, bold=True, color=INK2, anchor=MSO_ANCHOR.MIDDLE)
    hline(s, L + 0.6, 0.65, 0.3, MUTE, 0.75)
    tb(s, L + 1.02, 0.5, 6, 0.3, title, size=13, color=MUTE, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, R - 1.0, 0.47, 0.86, 0.36, 'JIL', size=18, bold=True, color=INK2, align=PP_ALIGN.RIGHT,
       anchor=MSO_ANCHOR.MIDDLE)
    rect(s, R - 0.11, 0.69, 0.075, 0.075, fill=ACC, shape=MSO_SHAPE.OVAL)
    tb(s, L, 1.1, W, 0.7, lead, size=30, bold=True, color=INK2, spacing=1.0)
    return s


def head(s, x, y, w, text, color=MUTE):
    """表の列見出し（灰色の小さな文字）"""
    tb(s, x, y, w, 0.3, text, size=14, bold=True, color=color)


def num(s, x, y, n, size=16):
    tb(s, x, y, 0.5, 0.35, f'{n:02d}', size=size, bold=True, color=ACC)


def arrow(s, x1, y, x2, color=ACC, lw=1.25):
    """右向きの矢印（線＋矢じり）。影は付けない"""
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y), Inches(x2), Inches(y))
    st = c._element.find(qn('p:style'))
    if st is not None:
        c._element.remove(st)
    c.line.color.rgb = color; c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    return c


def box(s, x, y, w, h, text, kind='node', size=16):
    """サイトツリーの箱。top＝墨の面、node＝墨の枠、ext＝青の枠（別タブで開く先）、label＝灰色の枠（押せない見出し）"""
    if kind == 'top':
        rect(s, x, y, w, h, fill=INK2)
        col = WHITE
    elif kind == 'label':
        rect(s, x, y, w, h, fill=WHITE, line=MUTE, lw=1.0)
        col = MUTE
    elif kind == 'ext':
        rect(s, x, y, w, h, fill=WHITE, line=ACC, lw=1.25)
        col = ACC
    else:
        rect(s, x, y, w, h, fill=WHITE, line=INK2, lw=1.0)
        col = INK2
    tb(s, x, y, w, h, text, size=size, bold=True, color=col, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
       spacing=1.0)


def note(s, text):
    tb(s, L, 6.95, W - 1.3, 0.3, text, size=12, color=MUTE)


prs = new_deck()
# テーマの標準の効果（影）も空にして、影の設定を0件にする
import re as _re
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
_th = prs.slide_master.part.part_related_by(RT.THEME)
_th._blob = _re.sub(r'<a:effectStyle>.*?</a:effectStyle>', '<a:effectStyle><a:effectLst/></a:effectStyle>',
                    _th.blob.decode('utf-8'), flags=_re.S).encode('utf-8')

# ---------- 1. 今のサイトの課題 ----------
s = slide(prs, 1, '今のサイトの課題', '会社の顔つきが見えにくい今のサイト')
y = Y0 + 0.1
hline(s, L, y, 5.2, INK2, 1.25)
for i, t in enumerate(['最初の画面は英字の標語と商品写真。\n何の会社か読み取りにくい',
                       '載っている事業は2つ。\nすでに提供中の採用支援が載っていない',
                       '働く人・組織の様子が伝わらない']):
    num(s, L, y + 0.28, i + 1)
    tb(s, L + 0.6, y + 0.25, 4.6, 0.9, t, size=18, color=INK2)
    y += 1.3
    hline(s, L, y, 5.2)
pic(s, IMG / 'corp_old_top.jpg', 6.3, Y0 + 0.1, w=6.28, border=True)
note(s, '出典：https://jilinc.co.jp/ （2026/10/23 撮影）')

# ---------- 2. リニューアルの目標 ----------
s = slide(prs, 2, 'リニューアルの目標', '「人が中心の、AIで動く会社」が伝わるサイトへ')
cw = (W - 0.8) / 3
goals = [('ひと目で\n「何の会社か」がわかる', None),
         ('3つの事業が\nそろって見える', 'プロダクト／新規事業開発支援／\n採用支援'),
         ('読んだ人が迷わず、\nお問い合わせまで\nたどり着ける', None)]
for i, (t, sub) in enumerate(goals):
    x = L + i * (cw + 0.4)
    hline(s, x, Y0, cw, INK2, 1.25)
    num(s, x, Y0 + 0.17, i + 1)
    tb(s, x, Y0 + 0.55, cw, 1.2, t, size=20, bold=True, color=INK2)
    if sub:
        tb(s, x, Y0 + 1.42, cw, 0.6, sub, size=14, color=MUTE)
y = 4.5
head(s, L, y, 4, '主な読み手')
head(s, L + 4.2, y, 6, '知りたいこと')
y += 0.38
hline(s, L, y, W, INK2, 1.25)
for who, want in [('発注を考える企業', '何を頼めるか、任せて大丈夫か'),
                  ('一緒に働きたい人', 'どんな考え方・どんな人の会社か'),
                  ('取引先・関係者', '会社の概要、役員')]:
    tb(s, L, y + 0.13, 4, 0.4, who, size=18, bold=True, color=INK2)
    tb(s, L + 4.2, y + 0.13, 7.5, 0.4, want, size=18, color=INK2)
    y += 0.56
    hline(s, L, y, W)

# ---------- 3. 伝えたいことの整理 ----------
s = slide(prs, 3, '伝えたいことの整理', '中心に人、周りにAI、外側に3つの事業')
y = Y0 + 0.25
hline(s, L, y, 6.5, INK2, 1.25)
for k, t in [('中心', '事業をつくる「人」の会社'),
             ('その周り', '組織のあらゆる仕事にAIを使う'),
             ('外側', '人とAIの力を、3つの事業に')]:
    tb(s, L, y + 0.3, 1.3, 0.4, k, size=16, bold=True, color=ACC)
    tb(s, L + 1.4, y + 0.26, 5.1, 0.45, t, size=20, bold=True, color=INK2)
    y += 1.0
    hline(s, L, y, 6.5)
tb(s, L + 1.4, y + 0.3, 5.1, 0.4, '3つの事業は、互いに知見を返し合う', size=18, color=MUTE)
pic(s, IMG / 'corp_new_top_shape_fig.png', R - 4.55, Y0 + 0.05, h=4.4)
note(s, '図：新しいトップページの「すべては、人から。」の欄')

# ---------- 4. 言葉の整理 ----------
s = slide(prs, 4, '言葉の整理', '今のサイトの言葉を受け継ぎ、新しい言葉を足す')
y = Y0 + 0.05
head(s, L, y, 2.6, '場所')
head(s, L + 2.8, y, 8, '言葉')
y += 0.38
hline(s, L, y, W, INK2, 1.25)
for where, words in [('最初の画面', '日本の可能性を、解き放つ。\nAI時代を生きる、私たちの手と意思で。'),
                     ('会社の一言', 'JILは、事業をつくる会社です。'),
                     ('私たちについて', '日本の現場力と、AIの進化。そのあいだを、人で繋ぐ。'),
                     ('会社の構造', 'すべては、人から。'),
                     ('人とAI', 'AIが手を増やし、人が意思を磨く。')]:
    h = 1.0 if '\n' in words else 0.62
    tb(s, L, y + 0.17, 2.6, 0.4, where, size=16, color=MUTE)
    tb(s, L + 2.8, y + 0.13, 9, h - 0.2, words, size=20, bold=True, color=INK2, spacing=1.2)
    y += h
    hline(s, L, y, W)
tb(s, L, y + 0.25, W, 0.4, [[('「JAPAN IS LOVELY.」', {'bold': True}), ('は社名の由来として残し、フッターに大きく置く', {})]],
   size=18, color=INK2)

# ---------- 5. 参考サイトの調査 ----------
s = slide(prs, 5, '参考サイトの調査', '国内外40社のコーポレートサイトを比較')
y = Y0 + 0.05
hline(s, L, y, 4.6, INK2, 1.25)
for cat, n, ex in [('AI企業（国内）', 7, 'PKSHA／Algomatic／LayerX／Sakana AI ほか'),
                   ('AI企業（海外）', 6, 'Anthropic／Cohere／Mistral ほか'),
                   ('スタートアップ', 8, 'SmartHR／10X／Ubie／UPSIDER ほか'),
                   ('制作・ブランディング（国内）', 8, 'Takram／Goodpatch／BASSDRUM ほか'),
                   ('制作・ブランディング（海外）', 5, 'Pentagram／Instrument ほか'),
                   ('ITコンサル・人材・組織', 6, 'Ridgelinez／Sun*／MIMIGURI ほか')]:
    tb(s, L, y + 0.1, 3.9, 0.35, cat, size=16, bold=True, color=INK2)
    tb(s, L, y + 0.42, 3.9, 0.3, ex, size=12, color=MUTE)
    tb(s, L + 3.7, y + 0.08, 0.9, 0.55, [[(str(n), {'size': 24, 'bold': True}), ('社', {'size': 13})]],
       color=INK2, align=PP_ALIGN.RIGHT, spacing=1.0)
    y += 0.73
    hline(s, L, y, 4.6)
pic(s, IMG / 'corp_refs_40.jpg', 5.75, Y0 + 0.05, w=6.83, border=True)
tb(s, 5.75, Y0 + 2.85, 6.83, 0.6, '40社の最初の画面（左上から、左の表の分野の順に並べている）',
   size=14, color=MUTE)
note(s, '出典：各社サイト（2026/10/23 撮影）')

# ---------- 6. 手本の選び方 ----------
s = slide(prs, 6, '手本の選び方', '決め手は「人の集まりに見えるか」')
head(s, L, Y0 - 0.15, 6, '手本にした4社')
iw, ih = 2.7, 2.7 * 0.625
for i, (f, name) in enumerate([('pksha', 'PKSHA Technology'), ('algomatic', 'Algomatic'),
                               ('layerx', 'LayerX'), ('smarthr', 'SmartHR')]):
    x = L + (i % 2) * (iw + 0.2)
    yy = Y0 + 0.2 + (i // 2) * (ih + 0.45)
    pic(s, IMG / f'corp_ref_{f}.jpg', x, yy, w=iw, h=ih, border=True)
    tb(s, x, yy + ih + 0.06, iw, 0.3, name, size=15, bold=True, color=INK2)
x = 6.95
head(s, x, Y0 - 0.15, 5, '手本から外した理由（例）')
y = Y0 + 0.2
hline(s, x, y, R - x, INK2, 1.25)
for why, ex in [('商品紹介のページに見える', 'ELYZA／PFN／Sakana AI'),
                ('製品が主役すぎる', 'Harvey／10X'),
                ('外資っぽい', 'Anthropic／Mistral／Runway'),
                ('暗い・素っ気ない', 'Linear／Visional')]:
    tb(s, x, y + 0.15, R - x, 0.4, why, size=18, bold=True, color=INK2)
    tb(s, x, y + 0.55, R - x, 0.3, ex, size=14, color=MUTE)
    y += 0.98
    hline(s, x, y, R - x)
note(s, '出典：pkshatech.com／algomatic.jp／layerx.co.jp／smarthr.co.jp（2026/10/23 撮影）')

# ---------- 7. 手本4社の共通点 ----------
s = slide(prs, 7, '手本4社の共通点', '4社に共通する6つの型と、JILでの取り入れ方')
y = Y0 + 0.05
c2 = L + 5.3
head(s, L, y, 4.5, '4社の共通点')
head(s, c2, y, 6, 'JILでの取り入れ方', ACC)
y += 0.38
hline(s, L, y, W, INK2, 1.25)
for a, b in [('会社を言い切る一文', '日本の可能性を、解き放つ。'),
             ('会社ならではのモチーフ', '水面に落ちる一滴（意思が広がる）'),
             ('中心から広がる構造図', '人 → AI → 3つの事業の図'),
             ('厚い「私たちについて」', '独立したページに（ミッション〜会社概要）'),
             ('明るい地に、ブランド色を1色', '白と墨に、青をひとつ'),
             ('条件をそろえた社員写真', '撮影は役員3名のみ。組織感は図と言葉で出す')]:
    tb(s, L, y + 0.15, 4.5, 0.4, a, size=18, color=INK2)
    arrow(s, L + 4.45, y + 0.33, L + 5.0)
    tb(s, c2, y + 0.15, R - c2, 0.4, b, size=18, bold=True, color=INK2)
    y += 0.66
    hline(s, L, y, W)

# ---------- 8. 検討の経緯 ----------
s = slide(prs, 8, '検討の経緯', '試作を見比べ、引き算で決めた')
y = Y0 + 0.05
xa, xb, xc = L, L + 2.4, L + 7.4
head(s, xa, y, 2.3, '検討したこと')
head(s, xb, y, 4.8, '試した案')
head(s, xc, y, 4.4, '決めたこと', ACC)
y += 0.38
hline(s, L, y, W, INK2, 1.25)
for a, b, c in [('最初の画面', '波紋の図／点の集まり／人の写真', '本物の水面に一滴が落ちる動き'),
                ('人の表し方', '写真／指紋／筆文字／シルエット／アイコン6種', '2色の人のアイコン'),
                ('人とAIの図', '掛け算の帯／2つの円の重なり', '円の重なり（透けて見える形）'),
                ('色', '青×橙／青×黄／青紫×桃', '青ひとつ。\n色は下層ページの冒頭で足す'),
                ('写真', '夜景／工場／職人の手 など', '人が写る写真\n（交差点・頂上の人たち）')]:
    h = 0.92 if '\n' in c else 0.66
    tb(s, xa, y + 0.15, 2.3, 0.4, a, size=18, bold=True, color=INK2)
    tb(s, xb, y + 0.17, 4.8, 0.4, b, size=16, color=MUTE)
    tb(s, xc, y + 0.15, R - xc, h - 0.2, c, size=18, bold=True, color=INK2, spacing=1.15)
    y += h
    hline(s, L, y, W)
tb(s, L, y + 0.2, W, 0.4, '24回の改版。毎回、画面を撮って手本と並べて確認', size=16, color=MUTE)

# ---------- 9. デザインの方針 ----------
s = slide(prs, 9, 'デザインの方針', '白と墨に、青をひとつ')
y = Y0 + 0.05
for name, hexv, col, use in [('墨', '#0E1013', INK2, '文字'), ('白', '#F1F2F4', PAPER, '地'),
                             ('青', '#2D44E0', ACC, '強調と「人」の印だけ')]:
    rect(s, L, y, 0.5, 0.5, fill=col, line=LINE if col == PAPER else None)
    tb(s, L + 0.7, y - 0.02, 1.2, 0.3, name, size=18, bold=True, color=INK2)
    tb(s, L + 0.7, y + 0.28, 1.6, 0.25, hexv, size=11, color=MUTE, name=MONO)
    tb(s, L + 2.0, y + 0.06, 2.8, 0.4, use, size=16, color=INK2)
    y += 0.64
y += 0.08
hline(s, L, y, 4.75, INK2, 1.25)
for t in ['動きは最初の画面の水面と、\n図ができあがる動きに絞る', '写真は「人が写るもの」だけ',
          '英字は Satoshi、\n日本語は Zen Kaku Gothic New']:
    h = 0.88 if '\n' in t else 0.58
    tb(s, L, y + 0.12, 4.75, h - 0.15, t, size=18, color=INK2, spacing=1.2)
    y += h
    hline(s, L, y, 4.75)
pic(s, IMG / 'corp_new_top_fv.png', 5.85, Y0 + 0.05, w=6.73, border=True)
tb(s, 5.85, Y0 + 4.3, 6.73, 0.3, '新しいトップページの最初の画面', size=12, color=MUTE)

# ---------- 10. サイトツリー ----------
s = slide(prs, 10, 'サイトツリー', 'トップ＋5ページ＋専用サイト1つ')
bh = 0.46
rows = {'about': 2.55, 'product': 3.15, 'svc': 3.75, 'member': 4.95, 'news': 5.5, 'contact': 6.05}
kids = [3.75, 4.35]
x1, x2, x3, x4 = L, 3.2, 6.15, 9.55
ty = 4.45
box(s, x1, ty - 0.28, 1.75, 0.56, 'トップ', 'top', 18)
trunk = x1 + 1.75 + 0.35
hline(s, x1 + 1.75, ty, 0.35, INK2, 1.0)
vline(s, trunk, rows['about'], rows['contact'] - rows['about'], INK2, 1.0)
for k, label_text in [('about', '私たちについて'), ('product', 'プロダクト'), ('svc', 'サービス'), ('member', 'メンバー'),
                      ('news', 'ニュース'), ('contact', 'お問い合わせ')]:
    hline(s, trunk, rows[k], x2 - trunk, INK2, 1.0)
    box(s, x2, rows[k] - bh / 2, 2.45, bh, label_text, 'label' if k == 'svc' else 'node')
tb(s, x2, rows['svc'] + bh / 2 + 0.04, 2.45, 0.3, '押せない見出し', size=12, color=MUTE, align=PP_ALIGN.CENTER)
ktrunk = x2 + 2.45 + 0.25
hline(s, x2 + 2.45, rows['svc'], 0.25, INK2, 1.0)
vline(s, ktrunk, kids[0], kids[1] - kids[0], INK2, 1.0)
for yk, t in zip(kids, ['新規事業開発支援', '採用支援']):
    hline(s, ktrunk, yk, x3 - ktrunk, INK2, 1.0)
    box(s, x3, yk - bh / 2, 2.55, bh, t)
for k, t in [('about', 'ミッション・ビジョン／行動指針／スタイル／会社概要'),
             ('member', 'ボードメンバー／チームメンバー'), ('news', 'トップの中の欄')]:
    tb(s, x3, rows[k] - 0.17, 6.4, 0.35, t, size=15, color=MUTE)
arrow(s, x3 + 2.55, kids[0], x4)
box(s, x4, kids[0] - bh / 2, R - x4, bh, '専用サイト（サービスページ）', 'ext', 15)
arrow(s, x2 + 2.45, rows['contact'], x4)
box(s, x4, rows['contact'] - bh / 2, R - x4, bh, 'お問い合わせフォーム', 'ext', 15)
tb(s, x4, rows['contact'] - 0.95, R - x4, 0.4, '青の枠：別のタブで開く', size=13, color=ACC, align=PP_ALIGN.RIGHT)
tb(s, L, 6.5, W, 0.35, 'メニューは「プロダクト」と「サービス」に分け、自社のプロダクトとお客様向けのサービスを見分けやすくした',
   size=16, color=INK2)
note(s, 'ページの住所は今の会社サイトと同じ形（/about/、/product/ など）')

# ---------- 11. トップページの流れ ----------
s = slide(prs, 11, 'トップページの流れ', '上から読むと「JILとは何か」がわかる順番')
y = Y0 + 0.0
xn, xs, xr = L, L + 0.55, L + 2.75
head(s, xs, y, 2.1, '欄')
head(s, xr, y, 3, '役割')
y += 0.38
hline(s, L, y, 5.55, INK2, 1.25)
for i, (sec, role) in enumerate([('最初の画面', '何の会社かを一言で'), ('私たちについて', '想いと、やっていること'),
                                 ('会社の構造', '人 → AI → 3つの事業'), ('3つの事業', '各事業ページへの入口'),
                                 ('人とAI', '働き方の考え方'), ('JILをつくる人たち', '役員3名'),
                                 ('ニュース', '会社の動き'), ('お問い合わせ', '採用・事業の2つの出口')]):
    num(s, xn, y + 0.1, i + 1, 14)
    tb(s, xs, y + 0.08, 2.2, 0.35, sec, size=16, bold=True, color=INK2)
    tb(s, xr, y + 0.08, 3.0, 0.35, role, size=16, color=INK2)
    y += 0.5
    hline(s, L, y, 5.55)
pic(s, IMG / 'corp_new_top_full3.jpg', 6.55, Y0 + 0.05, h=4.12, border=True)
tb(s, 6.55, Y0 + 4.25, 6, 0.3, 'トップページ全体（縦に長いページを3つに分けて並べたもの）', size=13, color=MUTE)

# ---------- 12. 各ページの役割 ----------
s = slide(prs, 12, '各ページの役割', '下層ページは同じ型で、役割を分ける')
y = Y0 - 0.15
xp, xn2, xr2, xc2 = L, L + 1.35, L + 3.85, L + 7.55
head(s, xn2, y, 2.3, 'ページ')
head(s, xr2, y, 3.5, '役割')
head(s, xc2, y, 4.2, '主な中身')
y += 0.36
hline(s, L, y, W, INK2, 1.25)
for f, name, role, what in [('about', '私たちについて', '考え方を伝える', '社名の由来／ミッション・ビジョン／\n行動指針／スタイル／会社概要'),
                            ('product', 'プロダクト', '自分たちでつくっている証拠', '4つのプロダクト／自社でつくる理由'),
                            ('bizdev', '新規事業開発支援', '強みだけ見せ、専用サイトへ', '見極める・つくる・動かす'),
                            ('recruit', '採用支援', '専用サイトがないので、\nここで全部', '考え方／よくあるお悩み／\n支援の流れ6工程／FAQ'),
                            ('member', 'メンバー', '人の顔が見える', '役員3名／チームの職種')]:
    pic(s, IMG / f'corp_new_{f}_fv.png', xp, y + 0.07, w=1.1, h=0.69, border=True)
    tb(s, xn2, y + 0.07, 2.4, 0.69, name, size=18, bold=True, color=INK2, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, xr2, y + 0.07, 3.6, 0.69, role, size=18, color=INK2, anchor=MSO_ANCHOR.MIDDLE, spacing=1.15)
    tb(s, xc2, y + 0.07, R - xc2, 0.69, what, size=14, color=MUTE, anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    y += 0.83
    hline(s, L, y, W)
note(s, '冒頭は共通の型：大きな英字の見出し（About. など）＋線の背景')

# ---------- 13. 導線の設計 ----------
s = slide(prs, 13, '導線の設計', 'どのページからも、次の行き先がある')
y = Y0 + 0.05
hline(s, L, y, 7.4, INK2, 1.25)
for i, t in enumerate(['全ページ共通のメニュー。中の欄まで直接行ける',
                       '下層ページの最後に「次のページへ」。5ページを一周',
                       'お問い合わせは、どこからでも1回押せばフォームへ']):
    num(s, L, y + 0.17, i + 1)
    tb(s, L + 0.6, y + 0.14, 6.8, 0.4, t, size=18, color=INK2)
    y += 0.7
    hline(s, L, y, 7.4)
pic(s, IMG / 'corp_new_about_end.png', R - 3.5, Y0 + 0.05, w=3.5, border=True)
tb(s, R - 3.5, Y0 + 2.27, 3.5, 0.3, '下層ページの最後の「次のページへ」', size=12, color=MUTE)
y = 5.0
head(s, L, y, 3, '読み手')
head(s, L + 2.6, y, 8, 'たどる道')
y += 0.36
hline(s, L, y, W, INK2, 1.25)
for who, path in [('発注を考える企業', ['トップ', 'サービス', '新規事業開発支援', '専用サイト', '相談']),
                  ('一緒に働きたい人', ['トップ', '私たちについて', 'メンバー', '採用のお問い合わせ']),
                  ('取引先・関係者', ['トップ', '私たちについて', '会社概要'])]:
    tb(s, L, y + 0.1, 2.5, 0.35, who, size=16, bold=True, color=INK2)
    runs = []
    for j, p in enumerate(path):
        if j:
            runs.append(('  →  ', {'color': ACC, 'bold': True}))
        runs.append((p, {}))
    tb(s, L + 2.6, y + 0.1, R - L - 2.6, 0.35, [runs], size=16, color=INK2)
    y += 0.5
    hline(s, L, y, W)

# ---------- 14. スマホと使いやすさ ----------
s = slide(prs, 14, 'スマホと使いやすさ', 'スマホでも、同じ内容を同じ質で')
y = Y0 + 0.05
hline(s, L, y, 6.9, INK2, 1.25)
for k, t in [('メニュー', 'どの位置からでも開く、全画面のメニュー'),
             ('速さ', '書体を共通化し、写真を圧縮して表示を軽く'),
             ('読みやすさ', 'iPhoneでも言葉の途中で改行しない。\n補足の文字の濃さを基準に合わせる'),
             ('共有', 'シェア画像・説明文を設定'),
             ('確認', '6ページ×スマホ・パソコンで、\n表示崩れ・リンク切れを自動チェック')]:
    h = 0.98 if '\n' in t else 0.76
    tb(s, L, y + 0.2, 1.5, 0.35, k, size=16, bold=True, color=ACC)
    tb(s, L + 1.6, y + 0.17, 5.3, h - 0.2, t, size=18, color=INK2, spacing=1.2)
    y += h
    hline(s, L, y, 6.9)
ph_h = 4.1
ph_w = ph_h * 780 / 1688
for i, (f, cap) in enumerate([('corp_new_sp_top', '最初の画面'), ('corp_new_sp_menu', '途中から開いたメニュー')]):
    x = R - ph_w - (1 - i) * (ph_w + 0.35)
    pic(s, IMG / f'{f}.png', x, Y0 + 0.05, w=ph_w, h=ph_h, border=True)
    tb(s, x, Y0 + ph_h + 0.12, ph_w + 0.3, 0.3, cap, size=12, color=MUTE)

prs.save(OUT)
print('保存:', OUT)

# 影の設定が無いことを確認
import zipfile, re
with zipfile.ZipFile(OUT) as z:
    n = sum(len(re.findall(r'outerShdw|innerShdw', z.read(f).decode('utf-8', 'ignore')))
            for f in z.namelist() if f.endswith('.xml'))
print('影の設定:', n, '件')

# 画像に書き出して確認（_preview/ はサービスページ企画書の分があるので、別のフォルダに出す）
import subprocess, shutil
prev = HERE / '_preview_corp'
shutil.rmtree(prev, ignore_errors=True)
subprocess.run(['powershell', '-NoProfile', '-Command', f'''
$app = New-Object -ComObject PowerPoint.Application
$p = $app.Presentations.Open("{OUT.resolve()}", $true, $false, $false)
$p.Export("{prev.resolve()}", "JPG", 1600, 900)
$p.Close()
$app.Quit()
'''], check=True)
print('画像:', prev)
